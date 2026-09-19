"""独立第三方校验：用 python-pptx 检查生成的 pptx（图片数量、尺寸、越界、备注）。

用法: python test/verify-latex.py test/out/latex-direct.pptx
"""
import sys
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1] if len(sys.argv) > 1 else "test/out/latex-direct.pptx"
prs = Presentation(path)
W, H = prs.slide_width, prs.slide_height
print(f"slide size: {Emu(W).inches:.2f}\" x {Emu(H).inches:.2f}\"  slides: {len(prs.slides)}")

total_pics = 0
total_shapes = 0
overflow = []
zero_size = []
notes = 0
text_chars = 0

for i, slide in enumerate(prs.slides, 1):
    pics = 0
    for shape in slide.shapes:
        total_shapes += 1
        if shape.shape_type == 13 or shape.__class__.__name__ == "Picture":  # PICTURE
            pics += 1
            total_pics += 1
            if shape.width <= 0 or shape.height <= 0:
                zero_size.append((i, shape.name))
            if (shape.left < -Emu(10000) or shape.top < -Emu(10000)
                    or shape.left + shape.width > W + Emu(10000)
                    or shape.top + shape.height > H + Emu(10000)):
                overflow.append((i, shape.name,
                                 round(Emu(shape.width).inches, 2), round(Emu(shape.height).inches, 2)))
        if shape.has_text_frame:
            text_chars += len(shape.text_frame.text.strip())
    if slide.has_notes_slide:
        notes += 1

print(f"pictures: {total_pics}")
print(f"shapes:   {total_shapes}")
print(f"notes slides: {notes}")
print(f"text chars:   {text_chars}")
print(f"zero-size pictures: {len(zero_size)} {zero_size[:5]}")
print(f"picture overflow:   {len(overflow)} {overflow[:5]}")

ok = total_pics > 0 and not zero_size and len(overflow) == 0
print("\n" + ("VALID: python-pptx 解析正常，图片尺寸与边界正常" if ok else "PROBLEM FOUND"))
sys.exit(0 if ok else 1)
