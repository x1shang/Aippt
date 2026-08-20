"""独立第三方校验：用 python-pptx 打开生成的 .pptx，输出每页文本与备注。"""
import sys
from pptx import Presentation
from pptx.util import Emu

path = sys.argv[1] if len(sys.argv) > 1 else "test/out/sample-tech-blue.pptx"

prs = Presentation(path)
print(f"slide size: {prs.slide_width} x {prs.slide_height} EMU "
      f"({Emu(prs.slide_width).inches:.2f}\" x {Emu(prs.slide_height).inches:.2f}\")")
print(f"slides: {len(prs.slides)}")

for i, slide in enumerate(prs.slides, 1):
    texts = []
    shapes = 0
    for shape in slide.shapes:
        shapes += 1
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                texts.append(t.replace("\n", " ⏎ ")[:90])
    notes = ""
    if slide.has_notes_slide:
        nt = slide.notes_slide.notes_text_frame.text.strip()
        if nt:
            notes = nt.replace("\n", " ⏎ ")
    print(f"\n--- Slide {i} ({shapes} shapes) ---")
    for t in texts:
        print(f"  [text] {t}")
    if notes:
        print(f"  [note] {notes}")

print("\nOK: python-pptx parsed successfully")
